#!/usr/bin/env python3
"""Extrai metadados, frequência e conteúdo dos arquivos xlsx exportados do diário
para um único JSON, que depois é lido pelo script de importação em Ruby.
"""
import json
import re
import sys
import glob
import os
import openpyxl

SRC_DIR = sys.argv[1] if len(sys.argv) > 1 else '.'
OUT_FILE = sys.argv[2] if len(sys.argv) > 2 else 'diarios.json'

MATRICULADO_RE = re.compile(r'\s*\(Matriculado em [\d/]+\)\s*$', re.IGNORECASE)


def clean_student_name(name):
    return MATRICULADO_RE.sub('', name or '').strip()


def parse_file(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    details = {}
    ws = wb['Detalhes do diário']
    for row in ws.iter_rows(values_only=True):
        if row and row[0] and row[1] is not None:
            key = str(row[0]).rstrip(':').strip()
            details[key] = row[1]

    freq_ws = wb['Frequências']
    freq_rows = list(freq_ws.iter_rows(values_only=True))
    header = freq_rows[0]
    date_columns = []
    for idx, col in enumerate(header):
        if isinstance(col, str) and re.match(r'^\d{2}/\d{2}/\d{4}$', col):
            date_columns.append((idx, col))

    students = []
    for row in freq_rows[1:]:
        if not row or row[1] is None:
            continue
        raw_name = str(row[1])
        marks = {}
        for idx, date in date_columns:
            mark = row[idx] if idx < len(row) else None
            if mark is None or mark == '':
                continue
            marks[date] = str(mark).strip()

        students.append({
            'raw_name': raw_name,
            'name': clean_student_name(raw_name),
            'birthdate': row[2] if len(row) > 2 else None,
            'marks': marks,
        })

    contents = []
    content_ws = wb['Conteúdos']
    for row in content_ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        date = row[0]
        text = row[1]
        if not text or not str(text).strip():
            continue
        contents.append({'date': str(date), 'content': str(text).strip()})

    return {
        'file': os.path.basename(path),
        'professor': details.get('Professor(es)'),
        'escola': details.get('Escola'),
        'curso': details.get('Curso'),
        'serie': details.get('Série'),
        'turma': details.get('Turma'),
        'turno': details.get('Turno'),
        'disciplina': details.get('Disciplina'),
        'periodo_letivo': details.get('Período letivo'),
        'data_inicio': str(details.get('Data de início')),
        'data_termino': str(details.get('Data de término')),
        'students': students,
        'contents': contents,
    }


def main():
    files = sorted(glob.glob(os.path.join(SRC_DIR, '*.xlsx')))
    result = []
    for f in files:
        print(f'Lendo {f}...', file=sys.stderr)
        result.append(parse_file(f))

    with open(OUT_FILE, 'w', encoding='utf-8') as fh:
        json.dump(result, fh, ensure_ascii=False, indent=2)

    print(f'{len(result)} diários exportados para {OUT_FILE}', file=sys.stderr)


if __name__ == '__main__':
    main()
